"""
Чтение Файла 1 «Отработанное время за месяц».

Преобразует Excel-таблицу в структуру Python для дальнейшей обработки.
"""

from datetime import date, timedelta
from openpyxl import load_workbook

from core.constants import (
    COL_FIO,
    COL_FIRST_DAY,
    MAX_DAYS,
    RESPONSIBLE_PREFIX,
    ROW_FIRST_PERSON,
)
from core.time_utils import parse_cell


def read_source(path: str, start_date: date):
    """
    Читает файл «Отработанное время за месяц».

    start_date — дата первой колонки с данными (обычно понедельник).
    Количество дней ограничено MAX_DAYS.

    Возвращает список словарей:
    [
        {
            'fio': 'Гайдукевич Людмила',
            'days': [{'date': date(...), 'income': ..., 'outcome': ...}, ...],
        },
        ...
    ]
    """
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    days_count = _count_days(sheet)
    dates = _build_dates(start_date, days_count)

    people = []
    row = ROW_FIRST_PERSON

    while row <= sheet.max_row:
        fio = _read_fio(sheet, row)
        if fio is None:
            row += 1
            continue
        if fio.lower().startswith(RESPONSIBLE_PREFIX):
            break

        days = _read_days(sheet, row, dates)
        people.append({"fio": fio, "days": days})
        row += 1

    workbook.close()
    return people


def _count_days(sheet) -> int:
    """
    Считает, сколько колонок с днями реально заполнено.

    Ограничение — MAX_DAYS. Пустые колонки в конце отбрасываются.
    """
    last_filled = COL_FIRST_DAY - 1
    for column in range(COL_FIRST_DAY, sheet.max_column + 1):
        if _column_has_data(sheet, column):
            last_filled = column

    days = last_filled - COL_FIRST_DAY + 1
    if days < 0:
        days = 0
    return min(days, MAX_DAYS)


def _column_has_data(sheet, column: int) -> bool:
    """Проверяет, есть ли в колонке хотя бы одно непустое значение."""
    for row in range(ROW_FIRST_PERSON, sheet.max_row + 1):
        value = sheet.cell(row=row, column=column).value
        if value not in (None, "", "-", "?", "—"):
            return True
    return False


def _build_dates(start_date: date, count: int):
    """Строит список последовательных дат от start_date."""
    return [start_date + timedelta(days=i) for i in range(count)]


def _read_fio(sheet, row: int):
    """
    Возвращает ФИО из строки или None, если строка пустая.

    Служебные строки (без ФИО) игнорируются.
    """
    raw = sheet.cell(row=row, column=COL_FIO).value
    if raw is None:
        return None
    text = str(raw).strip()
    return text or None


def _read_days(sheet, row: int, dates):
    """Собирает данные по каждому дню для одного человека."""
    days = []
    for index, day_date in enumerate(dates):
        column = COL_FIRST_DAY + index
        cell = sheet.cell(row=row, column=column).value
        income, outcome = parse_cell(cell)
        days.append(
            {
                "date": day_date,
                "income": income,
                "outcome": outcome,
                "weekday_index": day_date.weekday(),
            }
        )
    return days
