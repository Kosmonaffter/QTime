"""
Запись файла СК — сверка Болида с ЗУП.

Пишет только расхождения. ОК-дни пропускаются.
Каждый сотрудник — блок с итоговой строкой.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.sk.constants import (
    MINUTES_IN_HOUR,
    NO_DATA_IN_BOLID,
    NO_DATA_IN_ZUP,
    SK_FONT_LATE_BOLD,
    SK_FONT_LATE_COLOR,
    SK_HEADERS,
    SK_OUT_COL_DIFF,
    SK_OUT_COL_FACT,
    SK_OUT_COL_FIO,
    SK_OUT_COL_INCOME,
    SK_OUT_COL_OUTCOME,
    SK_TITLE,
    STATUS_GREEN,
    STATUS_GRAY,
    STATUS_RED,
    STATUS_YELLOW,
)
from core.sk.constants import NO_DATA_IN_BOLID, NO_DATA_IN_ZUP


# ---------- Стили ----------
BORDER_COLOR = "999999"
HEADER_FILL = "DCE6F1"
HEADER_FONT_SIZE = 10
BODY_FONT_SIZE = 10

ROW_FILL = {
    STATUS_RED: "FFD6D6",
    STATUS_YELLOW: "FFF0C2",
    STATUS_GRAY: "EDEDED",
}

TOTAL_FILL = "E0E0E0"
TITLE_FILL = "C9DAF8"

COLUMN_WIDTHS = {
    "A": 20,
    "B": 12,
    "C": 11,
    "D": 11,
    "E": 12,
    "F": 14,
}


def write_sk(people, out_path: str):
    """Сохраняет сверку в Excel."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Не удалось создать активный лист")
    sheet.title = SK_TITLE

    styles = _build_styles()
    row = 1

    for person in people:
        row = _write_person(sheet, row, person, styles)

    _apply_column_widths(sheet)
    workbook.save(out_path)
    workbook.close()


def _build_styles():
    """Возвращает словарь со стилями."""
    thin = Side(style="thin", color=BORDER_COLOR)
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center"),
        "header_font": Font(bold=True, size=HEADER_FONT_SIZE),
        "body_font": Font(size=BODY_FONT_SIZE),
        "title_font": Font(bold=True, size=HEADER_FONT_SIZE + 1),
    }


def _write_person(sheet, row: int, person, styles):
    """Пишет блок одного сотрудника. Возвращает следующую строку."""
    deviations = _collect_deviations(person)

    if not deviations and person["total_ok"]:
        return row

    row = _write_title(sheet, row, person, styles)
    row = _write_headers(sheet, row, styles)

    for day in deviations:
        row = _write_day(sheet, row, day, styles)

    row = _write_total(sheet, row, person, styles)
    return row + 1


def _write_title(sheet, row: int, person, styles):
    """Пишет ФИО сотрудника и его график."""
    title = person["fio"]
    if person["graph"]:
        title = f"{title} — {person['graph']}"
    if person["no_bolid"]:
        title = f"{title} ({NO_DATA_IN_BOLID})"
    if person["no_zup"]:
        title = f"{title} ({NO_DATA_IN_ZUP})"

    cell = sheet.cell(row=row, column=SK_OUT_COL_FIO, value=title)
    cell.font = styles["title_font"]
    cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
    cell.alignment = styles["left"]
    cell.border = styles["border"]
    return row + 1


def _write_headers(sheet, row: int, styles):
    """Пишет шапку таблицы расхождений."""
    for offset, title in enumerate(SK_HEADERS):
        column = SK_OUT_COL_FIO + offset
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = styles["header_font"]
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = styles["center"]
        cell.border = styles["border"]
    return row + 1


def _write_day(sheet, row: int, day, styles):
    """Пишет строку одного дня с расхождением."""
    values = (
        _format_date(day["date"]),
        day["income"] or "",
        day["outcome"] or "",
        day["fact_text"],
        _day_diff_text(day),
    )
    fill = ROW_FILL.get(day["status"])
    diff_font = _build_diff_font(day, styles)

    for offset, value in enumerate(values):
        column = SK_OUT_COL_FIO + offset
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = diff_font if column == SK_OUT_COL_DIFF else styles["body_font"]
        cell.alignment = styles["center"]
        cell.border = styles["border"]
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)

    return row + 1


def _build_diff_font(day, styles):
    """
    Строит шрифт для колонки «Расхождение».

    Если недоработка (diff < 0) — жирный красный.
    Иначе — обычный.
    """
    minutes = day.get("diff_minutes")
    if minutes is not None and minutes < 0:
        return Font(
            bold=SK_FONT_LATE_BOLD,
            color=SK_FONT_LATE_COLOR,
            size=BODY_FONT_SIZE,
        )
    return styles["body_font"]


def _write_total(sheet, row: int, person, styles):
    """Пишет итоговую строку по сотруднику."""
    fact = person["total_fact"]
    zup = person["total_zup"]
    diff = round(fact - zup, 2)
    sign = "+" if diff >= 0 else "\u2212"
    diff_text = _format_total_diff(abs(diff))

    text = (
        f"Общее время работы: "
        f"Факт {fact:.2f} | ЗУП {zup:.2f} | "
        f"Расхождение {sign}{diff_text}"
    )
    font = _build_total_font(diff, styles)
    cell = sheet.cell(row=row, column=SK_OUT_COL_FIO, value=text)
    cell.font = font
    cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
    cell.alignment = styles["left"]
    cell.border = styles["border"]

    for column in range(SK_OUT_COL_FIO + 1, SK_OUT_COL_FIO + len(SK_HEADERS)):
        empty = sheet.cell(row=row, column=column, value="")
        empty.fill = PatternFill("solid", fgColor=TOTAL_FILL)
        empty.border = styles["border"]

    return row + 1


def _format_total_diff(diff: float) -> str:
    """19.3 -> '19:18'. 0.5 -> '00:30'."""
    total_minutes = round(diff * MINUTES_IN_HOUR)
    hours, minutes = divmod(total_minutes, MINUTES_IN_HOUR)
    return f"{hours:02d}:{minutes:02d}"


def _build_total_font(diff: float, styles):
    """
    Строит шрифт для итога.

    Если недоработка (diff < 0) — жирный красный.
    Иначе — обычный заголовочный.
    """
    if diff < 0:
        return Font(
            bold=SK_FONT_LATE_BOLD,
            color=SK_FONT_LATE_COLOR,
            size=HEADER_FONT_SIZE,
        )
    return styles["header_font"]


def _collect_deviations(person):
    """Возвращает дни, которые нужно показать (не ОК)."""
    result = []
    for day in person["days"]:
        status = day["status"]
        if status == STATUS_GREEN:
            continue
        if _is_empty(day):
            continue
        result.append(day)
    return result


def _is_empty(day):
    """Пустой день — нет ни Болида, ни ЗУП, ни факта."""
    if day["income"] or day["outcome"] or day["zup"] is not None:
        return False
    return day["status"] == STATUS_GRAY


def _day_diff_text(day):
    """Текст расхождения для дня."""
    if day["diff_text"]:
        return day["diff_text"]
    if day["status"] == STATUS_YELLOW:
        if day["zup"] is None and (day["income"] or day["outcome"]):
            return NO_DATA_IN_ZUP
    return ""


def _format_date(day_date):
    """Дата в формате ДД.ММ.ГГГГ или пусто."""
    if day_date is None:
        return ""
    return day_date.strftime("%d.%m.%Y")


def _apply_column_widths(sheet):
    """Задаёт ширины колонок."""
    for column, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
