"""
Чтение листа «Месяц Год» из файла ЗУП.

Поддерживает оба формата: .xlsx (openpyxl) и .xls (xlrd).
Находит раздел «Склад» в колонке A и читает таблицу сотрудников
до первой пустой строки.
"""

import calendar
from pathlib import Path

from core.sk.constants import (
    ZUP_SECTION_MARKER,
    ZUP_FIO_COLUMN,
    ZUP_NAME_COLUMN,
    ZUP_GRAPH_COLUMN,
    ZUP_DAY_START_COLUMN,

)


MONTHS_RU = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)

XLS_EXTENSION = ".xls"
XLSX_EXTENSION = ".xlsx"


def read_zup(path: str, month: int, year: int):
    """..."""
    sheet_name = _build_sheet_name(month, year)
    grid = _read_grid(path, sheet_name)

    header_row = _find_section_header(grid)
    if header_row is None:
        raise ValueError("Не найден раздел «Склад» в файле ЗУП")

    return _read_people(grid, header_row, month, year)


def _build_sheet_name(month: int, year: int) -> str:
    """Месяц 8, год 2026 → 'Август 2026'."""
    return f"{MONTHS_RU[month - 1]} {year}"


# ---------- Чтение грида (общее представление) ----------

def _read_grid(path: str, sheet_name: str):
    """
    Читает лист в список списков значений.

    Возвращает grid — двумерный массив [row][col], 0-based.
    Автоматически определяет формат по расширению.
    """
    extension = Path(path).suffix.lower()
    if extension == XLS_EXTENSION:
        return _read_grid_xls(path, sheet_name)
    return _read_grid_xlsx(path, sheet_name)


def _read_grid_xlsx(path: str, sheet_name: str):
    """Читает .xlsx через openpyxl."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Лист «{sheet_name}» не найден в файле ЗУП")
    sheet = workbook[sheet_name]

    grid = []
    for row in sheet.iter_rows(values_only=True):
        grid.append(list(row))
    workbook.close()
    return grid


def _read_grid_xls(path: str, sheet_name: str):
    """Читает .xls через xlrd."""
    import xlrd

    workbook = xlrd.open_workbook(path)
    if sheet_name not in workbook.sheet_names():
        raise ValueError(f"Лист «{sheet_name}» не найден в файле ЗУП")
    sheet = workbook.sheet_by_name(sheet_name)

    grid = []
    for row_index in range(sheet.nrows):
        row = [sheet.cell_value(row_index, col) for col in range(sheet.ncols)]
        grid.append(row)
    return grid


# ---------- Логика (общая для обоих форматов) ----------

def _find_section_header(grid):
    """
    Ищет строку «Склад» в колонке A.

    Возвращает индекс строки-заголовка таблицы (0-based) или None.
    """
    fio_col = ZUP_FIO_COLUMN - 1
    for row_index, row in enumerate(grid):
        if fio_col >= len(row):
            continue
        value = row[fio_col]
        if value is None:
            continue
        if str(value).strip() == ZUP_SECTION_MARKER:
            return row_index + 1
    return None


def _read_people(grid, header_row, month, year):
    """Читает сотрудников, начиная со строки данных."""
    people = []
    row_index = header_row + 1

    while row_index < len(grid):
        fio = _read_fio(grid, row_index)
        if not fio:
            break
        graph = _read_graph(grid, row_index)
        days = _read_days(grid, row_index, month, year)
        people.append({"fio": fio, "graph": graph, "days": days})
        row_index += 1

    return people


def _read_fio(grid, row_index: int) -> str:
    """Возвращает 'Фамилия Имя' или пустую строку."""
    surname = _cell(grid, row_index, ZUP_FIO_COLUMN)
    name = _cell(grid, row_index, ZUP_NAME_COLUMN)
    if not surname:
        return ""
    return f"{surname} {name}".strip()


def _read_graph(grid, row_index: int) -> str:
    """Возвращает значение графика работы."""
    return _cell(grid, row_index, ZUP_GRAPH_COLUMN)


def _read_days(grid, row_index: int, month: int, year: int):
    """
    Читает часы по дням месяца.

    Число дней берём из календаря — не читаем лишние пустые
    колонки справа.
    """
    total_days = calendar.monthrange(year, month)[1]
    days = {}

    for day in range(1, total_days + 1):
        column = ZUP_DAY_START_COLUMN + (day - 1)
        value = _raw_cell(grid, row_index, column)
        days[day] = _to_hours(value)

    return days


def _to_hours(value):
    """Преобразует значение в число часов или None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _cell(grid, row_index: int, column: int) -> str:
    """Значение ячейки как строка (1-based column)."""
    value = _raw_cell(grid, row_index, column)
    if value is None:
        return ""
    return str(value).strip()


def _raw_cell(grid, row_index: int, column: int):
    """Значение ячейки из grid (1-based column)."""
    if row_index >= len(grid):
        return None
    row = grid[row_index]
    col_index = column - 1
    if col_index >= len(row):
        return None
    return row[col_index]
