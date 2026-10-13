"""
Чтение листа «Месяц Год» из файла Табеля.

Находит раздел «Склад» в колонке A, затем строку заголовка
(по слову «Фамилия»), затем динамически определяет колонку
первого дня месяца (по значению «1»).

Поддерживает оба формата: .xlsx (openpyxl) и .xls (xlrd).
"""

import calendar
from pathlib import Path

from core.sk.constants import (
    ERR_TABEL_DAY_MISMATCH,
    TABEL_FIO_COLUMN,
    TABEL_GRAPH_COLUMN,
    TABEL_HEADER_MARKER,
    TABEL_NAME_COLUMN,
    TABEL_SECTION_MARKER,
)


MONTHS_RU = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)

XLS_EXTENSION = ".xls"


def read_tabel(path: str, month: int, year: int):
    """
    Читает файл Табеля и возвращает список сотрудников раздела «Склад».

    path — путь к .xls или .xlsx файлу.
    month — номер месяца (1..12).
    year — год.
    """
    sheet_name = _build_sheet_name(month, year)
    grid = _read_grid(path, sheet_name)

    section_row = _find_section(grid)
    if section_row is None:
        raise ValueError("Не найден раздел «Склад» в файле Табеля")

    header_row = _find_header(grid, section_row)
    if header_row is None:
        raise ValueError(
            "Не найдена строка заголовка «Фамилия» в разделе «Склад»"
        )

    day_start = _find_day_start_column(grid[header_row])
    _validate_days(grid[header_row], day_start, month, year)

    return _read_people(grid, header_row, day_start, month, year)


def _build_sheet_name(month: int, year: int) -> str:
    """Месяц 9, год 2026 → 'Сентябрь 2026'."""
    return f"{MONTHS_RU[month - 1]} {year}"


# ---------- Чтение грида ----------

def _read_grid(path: str, sheet_name: str):
    """Читает лист в список списков значений."""
    extension = Path(path).suffix.lower()
    if extension == XLS_EXTENSION:
        return _read_grid_xls(path, sheet_name)
    return _read_grid_xlsx(path, sheet_name)


def _read_grid_xlsx(path: str, sheet_name: str):
    """Читает .xlsx через openpyxl."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    available = list(workbook.sheetnames)
    if sheet_name not in available:
        workbook.close()
        raise ValueError(
            f"Лист «{sheet_name}» не найден в файле Табеля.\n"
            f"Доступные листы: {', '.join(available)}"
        )
    sheet = workbook[sheet_name]

    grid = []
    for row in sheet.iter_rows(values_only=True):
        grid.append(list(row))
    workbook.close()
    return grid


def _read_grid_xls(path: str, sheet_name: str):
    """Читает .xls через xlrd."""
    import xlrd

    workbook = xlrd.open_workbook(path)
    available = workbook.sheet_names()
    if sheet_name not in available:
        raise ValueError(
            f"Лист «{sheet_name}» не найден в файле Табеля.\n"
            f"Доступные листы: {', '.join(available)}"
        )
    sheet = workbook.sheet_by_name(sheet_name)

    grid = []
    for row_index in range(sheet.nrows):
        row = [sheet.cell_value(row_index, col) for col in range(sheet.ncols)]
        grid.append(row)
    return grid


# ---------- Поиск структуры ----------

def _find_section(grid):
    """Ищет строку с маркером 'Склад' в колонке A."""
    fio_col = TABEL_FIO_COLUMN - 1
    for row_index, row in enumerate(grid):
        if fio_col >= len(row):
            continue
        value = row[fio_col]
        if value is None:
            continue
        if str(value).strip() == TABEL_SECTION_MARKER:
            return row_index
    return None


def _find_header(grid, section_row):
    """Ищет строку с 'Фамилия' ниже строки раздела."""
    fio_col = TABEL_FIO_COLUMN - 1
    marker = TABEL_HEADER_MARKER.lower()
    for row_index in range(section_row + 1, len(grid)):
        row = grid[row_index]
        if fio_col >= len(row):
            continue
        value = row[fio_col]
        if value is None:
            continue
        if str(value).strip().lower().startswith(marker):
            return row_index
    return None


def _find_day_start_column(header_row):
    """
    Находит колонку первого дня месяца в строке заголовка.

    Ищет первую ячейку, значение которой — '1'
    (без учёта текста после переноса строки, например '1\\nвт').
    """
    for col_index, value in enumerate(header_row):
        if value is None:
            continue
        if _first_line(value) == "1":
            return col_index
    return None


def _validate_days(header_row, day_start, month, year):
    """Проверяет, что дни идут 1, 2, 3, ..., N без пропусков."""
    if day_start is None:
        raise ValueError(
            "В табеле не найдена колонка первого дня месяца (1)"
        )

    total_days = calendar.monthrange(year, month)[1]
    for offset in range(total_days):
        col_index = day_start + offset
        if col_index >= len(header_row):
            raise ValueError(
                ERR_TABEL_DAY_MISMATCH.format(missing=offset + 1)
            )
        text = _first_line(header_row[col_index])
        expected = str(offset + 1)
        if text != expected:
            raise ValueError(
                ERR_TABEL_DAY_MISMATCH.format(missing=expected)
            )


def _first_line(value) -> str:
    """Возвращает первую строку значения (до переноса)."""
    if value is None:
        return ""
    text = str(value).strip()
    return text.split("\n")[0].strip()


# ---------- Чтение сотрудников ----------

def _read_people(grid, header_row, day_start, month, year):
    """Читает сотрудников, начиная со строки после заголовка."""
    people = []
    row_index = header_row + 1

    while row_index < len(grid):
        fio = _read_fio(grid, row_index)
        if not fio:
            break
        graph = _read_graph(grid, row_index)
        days = _read_days(grid, row_index, day_start, month, year)
        people.append({"fio": fio, "graph": graph, "days": days})
        row_index += 1

    return people


def _read_fio(grid, row_index: int) -> str:
    """Возвращает 'Фамилия Имя' или пустую строку."""
    surname = _cell(grid, row_index, TABEL_FIO_COLUMN)
    name = _cell(grid, row_index, TABEL_NAME_COLUMN)
    if not surname:
        return ""
    return f"{surname} {name}".strip()


def _read_graph(grid, row_index: int) -> str:
    """Возвращает значение графика работы."""
    return _cell(grid, row_index, TABEL_GRAPH_COLUMN)


def _read_days(grid, row_index: int, day_start: int, month: int, year: int):
    """Читает часы по дням месяца (число дней — из календаря)."""
    total_days = calendar.monthrange(year, month)[1]
    days = {}

    for day in range(1, total_days + 1):
        col_index = day_start + (day - 1)
        value = _raw_cell(grid, row_index, col_index + 1)
        days[day] = _to_hours(value)

    return days


def _to_hours(value):
    """Преобразует значение в число часов или None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _cell(grid, row_index: int, column: int) -> str:
    """Значение ячейки как строка (1-based column)."""
    value = _raw_cell(grid, row_index, column)
    if value is None:
        return ""
    return str(value).strip()


def _raw_cell(grid, row_index: int, column: int):
    """Значение ячейки из grid (1-based column)."""
    if row_index >= len(grid):
        return None
    row = grid[row_index]
    col_index = column - 1
    if col_index >= len(row):
        return None
    return row[col_index]
