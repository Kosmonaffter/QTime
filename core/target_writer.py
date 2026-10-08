"""
Запись Файла 2 — итогового табеля УРВ.

Структура блока сотрудника:
    A: ФИО        B: приход   C: уход   D: итого
    E: статус 1   F: статус 2   G: норма
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.color_utils import (
    header_color,
    pick_color,
    to_excel_color,
)
from core.constants import (
    ABSENT,
    BORDER_COLOR,
    COL_WIDTH_FIO,
    COL_WIDTH_NORM,
    COL_WIDTH_STATUS,
    COL_WIDTH_TIME,
    COL_WIDTH_TOTAL,
    DATE_FORMAT_OUT,
    HEADER_FILL_COLOR,
    HEADER_FIO,
    HEADER_INCOME,
    HEADER_OUTCOME,
    HEADER_TOTAL,
    OUT_COL_DATE,
    OUT_COL_INCOME,
    OUT_COL_NORM,
    OUT_COL_OUTCOME,
    OUT_COL_STATUS_1,
    OUT_COL_STATUS_2,
    OUT_COL_TOTAL,
    OUT_HEADER_NORM,
    STATUS_COLORS,
    STATUS_TEXT_COLORS,
    WEEKEND_INDEXES,
)
from core.schedule import evaluate_day
from core.time_utils import calc_total


HEADER_COUNT = 4    # ФИО, приход, уход, итого


def write_target(people, schedules, out_path: str):
    """
    Сохраняет итоговый УРВ.

    people — список сотрудников с днями.
    schedules — словарь {ФИО: Schedule} с настройками графиков.
    """
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Не удалось создать активный лист")

    sheet.title = "УРВ"
    styles = _build_styles()

    row = 1
    for index, person in enumerate(people):
        base = pick_color(index)
        schedule = schedules.get(person["fio"])

        _write_person_header(sheet, row, person["fio"], base, styles)
        row += 1

        for day in person["days"]:
            if _skip_empty_weekend(day):
                continue
            _write_day_row(sheet, row, day, base, schedule, styles)
            row += 1

        row += 1

    _apply_column_widths(sheet)
    workbook.save(out_path)
    workbook.close()


def _build_styles():
    """Возвращает словарь со стилями ячеек."""
    thin = Side(style="thin", color=BORDER_COLOR)
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "header_font": Font(bold=True),
        "center": Alignment(horizontal="center", vertical="center"),
    }


def _write_person_header(sheet, row, fio, base_color, styles):
    """Пишет шапку блока сотрудника."""
    fill = PatternFill(
        "solid",
        fgColor=to_excel_color(header_color(base_color)),
    )
    labels = (
        HEADER_FIO,
        HEADER_INCOME,
        HEADER_OUTCOME,
        HEADER_TOTAL,
        "", "", "",
    )
    for offset, label in enumerate(labels):
        column = OUT_COL_DATE + offset
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = styles["header_font"]
        cell.fill = fill
        cell.alignment = styles["center"]
        cell.border = styles["border"]
    sheet.cell(row=row, column=OUT_COL_DATE, value=fio)


def _write_day_row(sheet, row, day, base_color, schedule, styles):
    """Пишет строку одного дня с оценкой по графику."""
    income = day["income"]
    outcome = day["outcome"]
    total = calc_total(income, outcome)

    evaluation = _evaluate(income, outcome, total, schedule)
    row_fill = _pick_row_fill(base_color, evaluation.priority_status)

    values = (
        day["date"].strftime(DATE_FORMAT_OUT),
        _resolve_income_text(income, outcome),
        outcome or "",
        total,
        evaluation.status_late,
        evaluation.status_incomplete,
        evaluation.norm_text if schedule else "",
    )

    for offset, value in enumerate(values):
        column = OUT_COL_DATE + offset
        cell = sheet.cell(row=row, column=column, value=value)
        cell.fill = row_fill
        cell.alignment = styles["center"]
        cell.border = styles["border"]

    _tint_status_cells(sheet, row, evaluation)


def _evaluate(income, outcome, total, schedule):
    """Оценивает день, если есть график. Иначе — пустая оценка."""
    if schedule is None:
        return _empty_evaluation()
    return evaluate_day(income, outcome, total, schedule)


def _empty_evaluation():
    """Возвращает пустую оценку — для сотрудников без графика."""
    from core.schedule import DayEvaluation
    return DayEvaluation("", "", "", "")


def _pick_row_fill(base_color, priority_status):
    """
    Определяет заливку строки.

    Если есть статус — берём его цвет. Иначе — цвет сотрудника.
    """
    if priority_status and priority_status in STATUS_COLORS:
        color = STATUS_COLORS[priority_status]
        return PatternFill("solid", fgColor=to_excel_color(color))
    return PatternFill("solid", fgColor=to_excel_color(base_color))


def _tint_status_cells(sheet, row, evaluation):
    """Красит текст в статусных колонках в цвет статуса."""
    if evaluation.status_late:
        _set_text_color(
            sheet, row, OUT_COL_STATUS_1, evaluation.status_late
        )
    if evaluation.status_incomplete:
        _set_text_color(
            sheet, row, OUT_COL_STATUS_2, evaluation.status_incomplete
        )


def _set_text_color(sheet, row, column, text):
    """Устанавливает жирный цветной шрифт для ячейки статуса."""
    cell = sheet.cell(row=row, column=column)
    cell.value = text
    status_key = _detect_status_key(text)
    color = STATUS_TEXT_COLORS.get(status_key, "000000")
    cell.font = Font(bold=True, color=color)


def _detect_status_key(text: str) -> str:
    """Определяет статус по тексту ячейки."""
    for key in STATUS_TEXT_COLORS:
        if text.startswith(key):
            return key
    return ""


def _skip_empty_weekend(day) -> bool:
    """Пропускает пустые сб/вс."""
    if day["weekday_index"] not in WEEKEND_INDEXES:
        return False
    return not day["income"] and not day["outcome"]


def _resolve_income_text(income, outcome):
    """Определяет текст для колонки «приход»."""
    if not income and not outcome:
        return ABSENT
    return income or ""


def _apply_column_widths(sheet):
    """Задаёт ширины колонок."""
    widths = {
        "A": COL_WIDTH_FIO,
        "B": COL_WIDTH_TIME,
        "C": COL_WIDTH_TIME,
        "D": COL_WIDTH_TOTAL,
        "E": COL_WIDTH_STATUS,
        "F": COL_WIDTH_STATUS,
        "G": COL_WIDTH_NORM,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
